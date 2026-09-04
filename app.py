import streamlit as st
import pandas as pd
import io
import re
import unicodedata
import openpyxl
from openpyxl.styles import PatternFill

st.set_page_config(page_title="業務自動化ツール集", layout="wide")

# --- サイドバーで機能選択 ---
st.sidebar.title("📌 機能メニュー")
app_mode = st.sidebar.radio(
    "利用するツールを選択してください",
    ["1. 売上明細 照合・差分抽出", "2. 基幹システムCSV ➜ freee変換"]
)

# ==========================================
# ツール1: 売上明細 照合・差分抽出システム
# ==========================================
if app_mode == "1. 売上明細 照合・差分抽出":
    st.title("🚗 売上明細 照合・差分抽出システム")
    st.caption("`deals.csv` と `整備売上明細.csv` を照合し、不一致データを抽出します。")

    # --- ルール設定 ---
    st.sidebar.header("⚙️ 照合ルール設定")
    valid_kanjo_default = "整備, 点検整備, 車検, 鈑金, レンタル売上"
    valid_kanjo_input = st.sidebar.text_input("対象とする勘定科目 (カンマ区切り)", valid_kanjo_default)
    valid_kanjo = [x.strip() for x in valid_kanjo_input.split(",") if x.strip()]

    st.sidebar.markdown("---")
    st.sidebar.subheader("除外ルール設定")
    ex_tenken_pack = st.sidebar.checkbox("「点検整備」かつ「定期点検パック(オリジナル)」を除外", value=True)
    ex_denpyo_nashi = st.sidebar.checkbox("品目「伝票なし」を除外", value=True)
    ex_leaseup = st.sidebar.checkbox("「リースアップ」を含むデータを除外", value=True)

    ex_seibi_kinds_input = st.sidebar.text_input("整備側で除外する種別 (カンマ区切り)", "その他, 新車待ち代車")
    ex_seibi_kinds = [x.strip() for x in ex_seibi_kinds_input.split(",") if x.strip()]

    # ヘルパー関数
    def read_csv_safe(file):
        try:
            return pd.read_csv(file, encoding='cp932')
        except Exception:
            file.seek(0)
            return pd.read_csv(file, encoding='utf-8-sig')

    def clean_reg(text):
        if pd.isna(text): return ""
        text = unicodedata.normalize('NFKC', str(text))
        return text.replace(' ', '').replace(' ', '').replace('-', '')

    def normalize_type_deals(row):
        pimmoku = row.get('品目', '')
        kanjo = row.get('勘定科目', '')
        if pd.isna(pimmoku) or str(pimmoku).strip() == '':
            if kanjo == '車検': return '車検'
            return ""
        val = unicodedata.normalize('NFKC', str(pimmoku)).strip()
        if val == '12ヵ月点検': return '12ヶ月点検'
        if val == '納車前点検': return '納車'
        return val

    def normalize_type_seibi(val):
        if pd.isna(val): return ""
        val = unicodedata.normalize('NFKC', str(val)).strip()
        if val in ['整備', 'タイヤ交換']: return '一般整備'
        if val == '12ヵ月点検': return '12ヶ月点検'
        return val

    col1, col2 = st.columns(2)
    with col1:
        file_deals = st.file_uploader("1. deals.csv をアップロード", type=["csv"])
    with col2:
        file_seibi = st.file_uploader("2. 整備売上明細 CSV をアップロード", type=["csv"])

    if file_deals and file_seibi:
        try:
            df_deals_orig = read_csv_safe(file_deals)
            df_seibi_orig = read_csv_safe(file_seibi)

            req_deals_cols = ['備考', '勘定科目', '品目', '金額', '税額', '管理番号', '取引先', '発生日']
            req_seibi_cols = ['登録番号', '種別', '請求額（税抜）', '伝票No', '請求先名', '売上日']

            missing_deals = [c for c in req_deals_cols if c not in df_deals_orig.columns]
            missing_seibi = [c for c in req_seibi_cols if c not in df_seibi_orig.columns]

            if missing_deals:
                st.error(f"❌ deals.csv に必要な列が見つかりません: {', '.join(missing_deals)}")
                st.stop()
            if missing_seibi:
                st.error(f"❌ 整備売上明細CSV に必要な列が見つかりません: {', '.join(missing_seibi)}")
                st.stop()

            df_deals = df_deals_orig[df_deals_orig['勘定科目'].isin(valid_kanjo)].copy()

            exclude_deals_conditions = []
            if ex_tenken_pack:
                exclude_deals_conditions.append((df_deals['勘定科目'] == '点検整備') & (df_deals['品目'] == '定期点検パック(オリジナル)'))
            if ex_denpyo_nashi:
                exclude_deals_conditions.append(df_deals['品目'] == '伝票なし')
            if ex_leaseup:
                exclude_deals_conditions.append(df_deals['品目'].astype(str).str.contains('リースアップ', na=False))

            if exclude_deals_conditions:
                final_deals_mask = exclude_deals_conditions[0]
                for cond in exclude_deals_conditions[1:]:
                    final_deals_mask = final_deals_mask | cond
                df_deals = df_deals[~final_deals_mask].copy()

            df_seibi = df_seibi_orig.copy()
            exclude_seibi_conditions = []
            if ex_seibi_kinds:
                exclude_seibi_conditions.append(df_seibi['種別'].isin(ex_seibi_kinds))
            if ex_leaseup:
                exclude_seibi_conditions.append(df_seibi['種別'].astype(str).str.contains('リースアップ', na=False))

            if exclude_seibi_conditions:
                final_seibi_mask = exclude_seibi_conditions[0]
                for cond in exclude_seibi_conditions[1:]:
                    final_seibi_mask = final_seibi_mask | cond
                df_seibi = df_seibi[~final_seibi_mask].copy()

            df_deals['orig_index'] = df_deals.index
            df_seibi['orig_index'] = df_seibi.index

            df_deals['reg_clean'] = df_deals['備考'].apply(clean_reg)
            df_seibi['reg_clean'] = df_seibi['登録番号'].apply(clean_reg)

            df_deals['type_clean'] = df_deals.apply(normalize_type_deals, axis=1)
            df_seibi['type_clean'] = df_seibi['種別'].apply(normalize_type_seibi)

            df_deals['amount_clean'] = (df_deals['金額'] - df_deals['税額'].fillna(0)).round().astype(int)
            df_seibi['amount_clean'] = pd.to_numeric(df_seibi['請求額（税抜）'], errors='coerce').fillna(0).round().astype(int)

            df_deals['match_key'] = df_deals['type_clean'] + "_" + df_deals['reg_clean'] + "_" + df_deals['amount_clean'].astype(str)
            df_seibi['match_key'] = df_seibi['type_clean'] + "_" + df_seibi['reg_clean'] + "_" + df_seibi['amount_clean'].astype(str)

            seibi_matched_indices = set()
            deals_unmatched = []

            for idx, row in df_deals.iterrows():
                key = row['match_key']
                candidates = df_seibi[(df_seibi['match_key'] == key) & (~df_seibi['orig_index'].isin(seibi_matched_indices))]
                if len(candidates) > 0:
                    matched_idx = candidates.index[0]
                    seibi_matched_indices.add(df_seibi.loc[matched_idx, 'orig_index'])
                else:
                    deals_unmatched.append(row)

            deals_unmatched_df = pd.DataFrame(deals_unmatched)
            seibi_unmatched_df = df_seibi[~df_seibi['orig_index'].isin(seibi_matched_indices)].copy()

            out_deals = pd.DataFrame({
                '存在ファイル': 'deals.csvのみ',
                '車両登録番号': deals_unmatched_df['備考'] if not deals_unmatched_df.empty else [],
                '売上種別': deals_unmatched_df.apply(lambda r: '車検' if (pd.isna(r['品目']) and r['勘定科目']=='車検') else r['品目'], axis=1) if not deals_unmatched_df.empty else [],
                '金額（税抜）': deals_unmatched_df['amount_clean'] if not deals_unmatched_df.empty else [],
                '伝票番号/管理番号': deals_unmatched_df['管理番号'] if not deals_unmatched_df.empty else [],
                '取引先/請求先名': deals_unmatched_df['取引先'] if not deals_unmatched_df.empty else [],
                '勘定科目': deals_unmatched_df['勘定科目'] if not deals_unmatched_df.empty else [],
                '発生日/売上日': deals_unmatched_df['発生日'] if not deals_unmatched_df.empty else []
            })

            out_seibi = pd.DataFrame({
                '存在ファイル': '整備売上明細CSVのみ',
                '車両登録番号': seibi_unmatched_df['登録番号'],
                '売上種別': seibi_unmatched_df['種別'],
                '金額（税抜）': seibi_unmatched_df['amount_clean'],
                '伝票番号/管理番号': seibi_unmatched_df['伝票No'].astype(str),
                '取引先/請求先名': seibi_unmatched_df['請求先名'],
                '勘定科目': '-',
                '発生日/売上日': seibi_unmatched_df['売上日']
            })

            result_df = pd.concat([out_deals, out_seibi], ignore_index=True)

            st.markdown("---")
            st.subheader("📊 照合結果サマリー")
            
            m1, m2, m3 = st.columns(3)
            m1.metric("不一致データ総数", f"{len(result_df)} 件")
            m2.metric("deals.csv のみ存在", f"{len(out_deals)} 件")
            m3.metric("整備売上明細CSV のみ存在", f"{len(out_seibi)} 件")

            if not result_df.empty:
                st.dataframe(result_df, use_container_width=True)
                csv_data = result_df.to_csv(index=False, encoding='utf-8-sig')
                st.download_button(
                    label="📥 不一致データ一覧（CSV）をダウンロード",
                    data=csv_data,
                    file_name="不一致データ一覧.csv",
                    mime="text/csv"
                )
            else:
                st.success("🎉 すべてのデータが一致しました！不一致データはありません。")

        except Exception as e:
            st.error(f"⚠️ 処理中にエラーが発生しました: {e}")

# ==========================================
# ツール2: 基幹システムCSV ➜ freee変換
# ==========================================
elif app_mode == "2. 基幹システムCSV ➜ freee変換":
    st.title("📄 基幹システムCSV ➜ freeeインポート変換アプリ")

    uploaded_file = st.file_uploader("整備売上明細などのCSVファイルをアップロードしてください", type="csv")

    def convert_wareki_to_seireki(date_str):
        era_map = {'R': 2018, 'H': 1988, 'S': 1925}
        match = re.match(r'([RHS])(\d+)\.(\d+)\.(\d+)', str(date_str))
        if match:
            era, year, month, day = match.groups()
            seireki_year = era_map[era] + int(year)
            return f"{seireki_year}/{int(month):02d}/{int(day):02d}"
        return date_str

    mapping_dict = {
        "12ヵ月点検": ("12ヶ月点検", "点検整備", "整備"),
        "車検": ("", "車検", "整備"),
        "オイル交換": ("オイル交換", "点検整備", "整備"),
        "6ヶ月点検": ("6ヶ月点検", "点検整備", "整備"),
        "3ヵ月点検": ("3ヵ月点検", "点検整備", "整備"),
        "1ヶ月点検": ("1ヶ月点検", "点検整備", "整備"),
        "整備": ("一般整備", "整備", "整備"),
        "リコール": ("リコール", "整備", "整備"),
        "用品取付": ("用品取付", "整備", "物販"),
        "タイヤ交換": ("一般整備", "整備", "整備"),
        "鈑金": ("鈑金", "鈑金", "鈑金"),
        "鈑金(保険)": ("鈑金(保険)", "鈑金", "鈑金"),
        "鈑金(業者)": ("鈑金(業者)", "鈑金", "鈑金"),
        "レンタカー": ("レンタカー", "レンタル売上", "レンタル"),
        "納車": ("納車前点検", "点検整備", "整備"),
        "新車待ち代車": ("代車使用料", "雑収入", ""),
        "手数料等": ("事務代行手数料", "雑収入", ""),
    }

    def split_legal_cost(total_legal_cost):
        if total_legal_cost == 0: return 0, 0
        jibaiseki_patterns = [9960,8190,10150,10620,10820,10700,10900]
        stamp_fee = 1850
        weight_tax_patterns = [
            5000,6600,7500,8200,10000,12300,15000,16400,18900,20000,24600,34200,37800,41000,45600,49200,50400
        ]
        candidates = []
        for jibai in jibaiseki_patterns:
            for weight in weight_tax_patterns:
                if jibai + weight + stamp_fee == total_legal_cost:
                    candidates.append((weight + stamp_fee, jibai))
        for weight in weight_tax_patterns:
            if weight + stamp_fee == total_legal_cost:
                candidates.append((weight + stamp_fee, 0))

        if len(candidates) == 1:
            return candidates[0]
        return None, None

    if uploaded_file is not None:
        try:
            df = pd.read_csv(uploaded_file, encoding="cp932")
            freee_rows = []
            alert_rows = []

            for _, row in df.iterrows():
                sale_date = convert_wareki_to_seireki(row["売上日"])
                partner = row["請求先名"]
                memo = row["登録番号"]

                if row["種別"] == "車検":
                    confirm_flag = 0
                    force_jibaiseki_row = False
                    shaken_sales = (row["小計"] + row["消費税"]) - row["メンテパック合計（税込）"]
                    legal_total = row["諸費用（非課税）"]
                    weight_stamp, jibaiseki = split_legal_cost(legal_total)

                    if weight_stamp is None:
                        st.warning(f"⚠️ 諸費用判定不可: {legal_total}円 / {partner}")
                        weight_stamp = legal_total
                        jibaiseki = 0
                        force_jibaiseki_row = True
                        confirm_flag = 1
                        alert_rows.append(len(freee_rows) + 3)

                    freee_rows.append({
                        "収支区分":"収入", "発生日": sale_date, "取引先": partner, "品目": "",
                        "金額": shaken_sales, "勘定科目": "車検", "税区分": "課税売上10%", "メモタグ": "",
                        "部門": "整備", "備考": memo, "確認要否": confirm_flag
                    })
                    if weight_stamp:
                        freee_rows.append({
                            "収支区分":"", "発生日": "", "取引先": partner, "品目": "重量税",
                            "金額": weight_stamp, "勘定科目": "課税対象外売上", "税区分": "不課税", "メモタグ": "",
                            "部門": "整備", "備考": memo, "確認要否": confirm_flag
                        })
                    if jibaiseki or force_jibaiseki_row:
                        freee_rows.append({
                            "収支区分":"", "発生日": "", "取引先": partner, "品目": "自賠責保険",
                            "金額": jibaiseki, "勘定科目": "課税対象外売上", "税区分": "不課税", "メモタグ": "",
                            "部門": "整備", "備考": memo, "確認要否": confirm_flag
                        })

                else:
                    mapping = mapping_dict.get(row["種別"])
                    confirm_flag = 1 if mapping is None else 0
                    if mapping is None: mapping = ("","","")

                    freee_rows.append({
                        "収支区分":"収入", "発生日": sale_date, "取引先": partner, "品目": mapping[0],
                        "金額": (row["小計"] + row["消費税"]) - row["メンテパック合計（税込）"],
                        "勘定科目": mapping[1], "税区分": "課税売上10%", "メモタグ": "",
                        "部門": mapping[2], "備考": memo, "確認要否": confirm_flag
                    })

            freee_df = pd.DataFrame(freee_rows)
            display_df = freee_df.copy()
            display_df["確認要否"] = display_df["確認要否"].apply(lambda x: "要確認" if x == 1 else "")

            st.write("✅ freeeインポート形式のプレビュー")
            st.dataframe(display_df, use_container_width=True)

            col_down1, col_down2 = st.columns(2)

            # Excelダウンロード処理
            excel_buffer = io.BytesIO()
            with pd.ExcelWriter(excel_buffer, engine="openpyxl") as writer:
                freee_df.to_excel(writer, index=False, sheet_name="freee取引")

            excel_buffer.seek(0)
            wb = openpyxl.load_workbook(excel_buffer)
            ws = wb.active
            yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

            for row_num in alert_rows:
                ws[f"E{row_num}"].fill = yellow_fill

            output_stream = io.BytesIO()
            wb.save(output_stream)
            output_stream.seek(0)

            with col_down1:
                st.download_button(
                    "📥 Excelでダウンロード",
                    data=output_stream.getvalue(),
                    file_name="freee_import.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True
                )

            # CSVダウンロード処理
            csv = freee_df.to_csv(index=False).encode('cp932')
            with col_down2:
                st.download_button(
                    "📥 freee用CSVをダウンロード",
                    csv,
                    "freee_import.csv",
                    "text/csv",
                    use_container_width=True
                )

        except Exception as e:
            st.error(f"⚠️ 変換処理中にエラーが発生しました: {e}")
